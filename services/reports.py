"""Spreadsheet builders for the two teacher downloads.

Two distinct reports, deliberately kept apart:

* `build_detailed_report` - one row per student answer for a single assignment,
  filtered to one class. Carries the response, the AI's verdict, its one-line
  key point and the marking provenance. This is the diagnostic sheet you read
  when a mark needs justifying or correcting.

* `build_marksheet` - one row per student, one column per assignment, across a
  whole class. No answers or reasoning; just the register.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from extensions import db
from models import Log, Question, User

# Matches the app's paper palette so exported files do not look bolted on.
HEADER_FILL = PatternFill('solid', fgColor='F6F4EF')
HEADER_FONT = Font(bold=True, color='1A1A17')
TITLE_FONT = Font(bold=True, size=13, color='1A1A17')
MUTED_FONT = Font(size=9, color='6B675E')
THIN = Side(style='thin', color='E6E2D9')
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(ws, row_idx, count):
    for col in range(1, count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = CELL_BORDER
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=row_idx + 1, column=1)


def _autosize(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _short_title(title: str, used: set) -> str:
    """'DA1 - Processes and Threads' -> 'DA1', kept unique across columns."""
    short = (title or 'Set').split(' - ')[0].strip() or 'Set'
    if short not in used:
        used.add(short)
        return short
    suffix = 2
    while f'{short} ({suffix})' in used:
        suffix += 1
    unique = f'{short} ({suffix})'
    used.add(unique)
    return unique


def _pct(score, maximum):
    if not maximum:
        return ''
    return round((score / maximum) * 100, 1)


def build_detailed_report(set_obj, class_group=None):
    """Per-answer detail for one assignment, optionally limited to one class."""
    query = db.session.query(Log, Question, User).join(
        User, Log.student_id == User.id
    ).join(
        Question, Log.question_id == Question.id
    ).filter(Question.set_id == set_obj.id)

    if class_group is not None:
        member_ids = [member.id for member in class_group.members]
        if not member_ids:
            query = query.filter(False)  # an empty class yields an empty report
        else:
            query = query.filter(User.id.in_(member_ids))

    rows = query.order_by(User.name, Log.question_id, Log.timestamp).all()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Detailed marks'

    headers = ['Student', 'Email', 'Class', 'Question', 'Student response', 'Key point',
               'Score', 'Max', '%', 'AI validation', 'Source', 'Confidence',
               'Late', 'Manually graded', 'Teacher comment', 'Submitted at']

    ws['A1'] = f'{set_obj.title} - detailed marks'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = f"Class: {class_group.name if class_group else 'All classes'}"
    ws['A2'].font = MUTED_FONT

    header_row = 4
    for col, name in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=name)
    _style_header(ws, header_row, len(headers))

    for offset, (log, question, student) in enumerate(rows):
        classes = ', '.join(sorted(g.name for g in student.class_groups)) or '-'
        row = header_row + 1 + offset
        values = [
            student.name,
            student.email,
            classes,
            question.question_text,
            log.answer,
            log.key_point or '',
            log.score,
            log.max_score,
            _pct(log.score, log.max_score),
            log.ai_response or '',
            (log.source_used or '').replace('_', ' '),
            log.confidence,
            'Yes' if log.submitted_late else '',
            'Yes' if log.manually_graded else '',
            log.teacher_comment or '',
            log.timestamp.strftime('%Y-%m-%d %H:%M') if log.timestamp else '',
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=col in (4, 5, 6, 10))
        ws.cell(row=row, column=7).number_format = '0.##'
        ws.cell(row=row, column=12).number_format = '0.00'

    _autosize(ws, [18, 26, 10, 42, 46, 30, 7, 6, 7, 46, 14, 11, 7, 16, 24, 17])
    return wb


def build_marksheet(class_group, sets):
    """One row per student in the class, one column per assignment.

    Uses each student's latest attempt on each assignment so a retake replaces
    the earlier mark rather than adding to it. A blank cell means the student
    never attempted that assignment - it is not the same as a zero.
    """
    sets = list(sets)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Class mark sheet'

    used_titles = set()
    columns = [(set_obj, _short_title(set_obj.title, used_titles)) for set_obj in sets]

    ws['A1'] = f'{class_group.name} - mark sheet'
    ws['A1'].font = TITLE_FONT
    ws['A2'] = 'Latest attempt per assignment. Blank means not attempted.'
    ws['A2'].font = MUTED_FONT

    headers = ['Student', 'Email', 'Class'] + [name for _, name in columns] + ['Total', 'Percentage']
    header_row = 4
    for col, name in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=name)
    _style_header(ws, header_row, len(headers))

    members = sorted(class_group.members, key=lambda user: user.name)

    if sets and members:
        logs = db.session.query(Log, Question).join(
            Question, Log.question_id == Question.id
        ).filter(
            Question.set_id.in_([set_obj.id for set_obj in sets]),
            Log.student_id.in_([member.id for member in members]),
        ).all()
    else:
        logs = []

    # (student_id, set_id) -> {attempt_no: [score, max]}
    marks = {}
    for log, question in logs:
        key = (log.student_id, question.set_id)
        attempt = log.attempt_no or 1
        bucket = marks.setdefault(key, {}).setdefault(attempt, [0.0, 0])
        bucket[0] += log.score or 0.0
        bucket[1] += log.max_score or 0

    for offset, student in enumerate(members):
        row = header_row + 1 + offset
        classes = ', '.join(sorted(g.name for g in student.class_groups)) or '-'
        values = [student.name, student.email, classes]

        total_score, total_max = 0.0, 0
        for set_obj, _ in columns:
            per_attempt = marks.get((student.id, set_obj.id))
            if not per_attempt:
                values.append('')
                continue
            latest = max(per_attempt)  # latest attempt wins
            score, maximum = per_attempt[latest]
            values.append(round(score, 2))
            total_score += score
            total_max += maximum

        values.append(round(total_score, 2))
        values.append(_pct(total_score, total_max))

        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical='top')
            if col > 3 and isinstance(value, (int, float)):
                cell.number_format = '0.##'
        ws.cell(row=row, column=len(values)).number_format = '0.0'

    _autosize(ws, [20, 28, 10] + [12] * len(columns) + [10, 12])
    return wb


def workbook_bytes(workbook) -> BytesIO:
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream
